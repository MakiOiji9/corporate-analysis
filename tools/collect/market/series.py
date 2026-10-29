"""Time-series providers for collect-market-info.

Each provider has a pure parser (unit-tested with fixed payloads) and a thin fetcher.
Dates are normalised to YYYY-MM-DD: monthly = first day of month, quarterly = first
day of the quarter's first month, annual = Jan 1, daily/weekly = the date itself.

Providers (catalog "provider" field):
  boj        日銀 時系列統計 API（キー不要）           params: db, code
  dashboard  e-Stat 統計ダッシュボード API（キー不要）  params: indicator, cycle, seasonal
  mof_jgb    財務省 国債金利 CSV（キー不要）           params: tenor
  fred       FRED API（FRED_API_KEY 必須）             params: id
  cftc       CFTC COT (Socrata, キー不要)              params: contract
  jpx_flows  JPX 投資部門別売買状況 週次 xlsx（キー不要） params: investor
  yahoo      Yahoo Finance chart API（非公式・キー不要。業種 ETF の日次終値） params: symbol
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re
from dataclasses import dataclass
from typing import Callable

BOJ_API = "https://www.stat-search.boj.or.jp/api/v1/getDataCode"
DASHBOARD_API = "https://dashboard.e-stat.go.jp/api/1.0/Json/getData"
MOF_JGB_URLS = (
    "https://www.mof.go.jp/jgbs/reference/interest_rate/data/jgbcm_all.csv",
    "https://www.mof.go.jp/jgbs/reference/interest_rate/jgbcm.csv",
)
FRED_API = "https://api.stlouisfed.org/fred/series/observations"
CFTC_API = "https://publicreporting.cftc.gov/resource/6dca-aqww.json"
YAHOO_CHART = "https://query1.finance.yahoo.com/v8/finance/chart/{symbol}"  # unofficial, no key
JPX_FLOWS_PAGE = "https://www.jpx.co.jp/markets/statistics-equities/investor-type/index.html"
JPX_ORIGIN = "https://www.jpx.co.jp"
JPX_WEEKLY_FILE = re.compile(r'href="([^"]*stock_1_w_(\d{8})_(\d{8})\.xlsx)"')
ERA_BASE = {"R": 2018, "H": 1988, "S": 1925}
THOUSAND_YEN_PER_OKU = 100_000  # JPX reports 千円; raw tables use 億円


class ProviderError(Exception):
    """A series could not be fetched; the message is written to status.md as the reason."""


@dataclass(frozen=True)
class Obs:
    series_id: str
    date: str
    value: float


def _num(raw: object) -> float | None:
    if raw is None:
        return None
    text = str(raw).strip().replace(",", "")
    if text in ("", ".", "-", "－", "NaN", "null", "***"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


# ---------------------------------------------------------------- BOJ

def boj_date(raw: object, frequency: str) -> str:
    text = str(raw)
    if len(text) == 8:
        return f"{text[:4]}-{text[4:6]}-{text[6:]}"
    if len(text) == 6 and frequency.upper().startswith("QUARTER"):
        quarter = int(text[4:])
        return f"{text[:4]}-{(quarter - 1) * 3 + 1:02d}-01"
    if len(text) == 6 and frequency.upper().startswith("SEMI"):
        half = int(text[4:])
        return f"{text[:4]}-{(half - 1) * 6 + 1:02d}-01"
    if len(text) == 6:
        return f"{text[:4]}-{text[4:]}-01"
    if len(text) == 4:
        return f"{text}-01-01"
    raise ValueError(f"unknown BOJ date {raw!r} ({frequency})")


def parse_boj(payload: dict, series_id: str) -> list[Obs]:
    if payload.get("STATUS") != 200:
        raise ProviderError(f"BOJ API: {payload.get('MESSAGE', 'unknown error')}")
    out = []
    for item in payload.get("RESULTSET") or []:
        values = item.get("VALUES") or {}
        freq = item.get("FREQUENCY", "")
        for date, value in zip(values.get("SURVEY_DATES") or [], values.get("VALUES") or []):
            number = _num(value)
            if number is not None:
                out.append(Obs(series_id, boj_date(date, freq), number))
    return out


# ---------------------------------------------------------------- e-Stat dashboard

def dashboard_date(raw: str) -> str:
    year, rest = raw[:4], raw[4:]
    if rest in ("CY00", "FY00"):
        return f"{year}-01-01"
    if rest.endswith("Q00"):
        return f"{year}-{(int(rest[0]) - 1) * 3 + 1:02d}-01"
    return f"{year}-{rest[:2]}-01"


def parse_dashboard(payload: dict, series_id: str, cycle: str, seasonal: str | None) -> list[Obs]:
    stats = payload.get("GET_STATS", {})
    if stats.get("RESULT", {}).get("status") not in ("0", "1"):
        raise ProviderError(f"dashboard API: {stats.get('RESULT', {}).get('errorMsg', 'error')}")
    objs = stats.get("STATISTICAL_DATA", {}).get("DATA_INF", {}).get("DATA_OBJ") or []
    objs = objs if isinstance(objs, list) else [objs]
    picked: dict[str, float] = {}
    for obj in objs:
        v = obj.get("VALUE", {})
        if v.get("@cycle") != cycle or v.get("@regionCode", "00000") != "00000":
            continue
        if seasonal is not None and v.get("@isSeasonal") != seasonal:
            continue
        number = _num(v.get("$"))
        if number is not None:
            picked[dashboard_date(v["@time"])] = number  # later stats override duplicates
    return [Obs(series_id, d, picked[d]) for d in sorted(picked)]


# ---------------------------------------------------------------- MOF JGB

def wareki_date(raw: str) -> str | None:
    m = re.fullmatch(r"([RHS])(\d+)\.(\d+)\.(\d+)", raw.strip())
    if not m:
        return None
    year = ERA_BASE[m.group(1)] + int(m.group(2))
    return f"{year}-{int(m.group(3)):02d}-{int(m.group(4)):02d}"


def parse_mof_jgb(text: str, series_id: str, tenor: str) -> list[Obs]:
    rows = list(csv.reader(io.StringIO(text)))
    header_at = next((i for i, r in enumerate(rows) if r and r[0].strip() == "基準日"), None)
    if header_at is None or tenor not in rows[header_at]:
        raise ProviderError(f"MOF JGB CSV: column {tenor} not found")
    col = rows[header_at].index(tenor)
    out = []
    for row in rows[header_at + 1:]:
        date = wareki_date(row[0]) if row else None
        number = _num(row[col]) if date and len(row) > col else None
        if date and number is not None:
            out.append(Obs(series_id, date, number))
    return out


# ---------------------------------------------------------------- FRED / CFTC

def parse_fred(payload: dict, series_id: str) -> list[Obs]:
    if "observations" not in payload:
        raise ProviderError(f"FRED API: {payload.get('error_message', 'no observations')}")
    return [Obs(series_id, o["date"], n) for o in payload["observations"]
            if (n := _num(o.get("value"))) is not None]


def parse_yahoo(payload: dict, series_id: str) -> list[Obs]:
    """Daily closes from the Yahoo Finance chart API; dates are the exchange's local trading days."""
    chart = payload.get("chart") or {}
    if chart.get("error") or not chart.get("result"):
        raise ProviderError(f"Yahoo chart API: {(chart.get('error') or {}).get('description', 'no result')}")
    result = chart["result"][0]
    offset = int((result.get("meta") or {}).get("gmtoffset") or 0)
    closes = ((result.get("indicators") or {}).get("quote") or [{}])[0].get("close") or []
    out: dict[str, Obs] = {}
    for ts, close in zip(result.get("timestamp") or [], closes):
        number = _num(close)
        if number is not None:
            day = dt.datetime.fromtimestamp(int(ts) + offset, tz=dt.timezone.utc).date().isoformat()
            out[day] = Obs(series_id, day, number)  # a later bar of the same day wins
    return [out[d] for d in sorted(out)]


def epoch(date: str) -> int:
    return int(dt.datetime.fromisoformat(date).replace(tzinfo=dt.timezone.utc).timestamp())


def parse_cftc(rows: list[dict], series_id: str) -> list[Obs]:
    out = []
    for r in rows:
        long_, short = _num(r.get("noncomm_positions_long_all")), _num(r.get("noncomm_positions_short_all"))
        if long_ is not None and short is not None:
            out.append(Obs(series_id, r["report_date_as_yyyy_mm_dd"][:10], long_ - short))
    return sorted(out, key=lambda o: o.date)


# ---------------------------------------------------------------- JPX investor flows

def _cell_text(ws, row: int, col: int) -> str:
    """Value of a (possibly merged) cell, as text."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return str(ws.cell(rng.min_row, rng.min_col).value or "").strip()
    return str(ws.cell(row, col).value or "").strip()


def parse_jpx_flows(xlsx: bytes, week_end: str, wanted: dict[str, str],
                    market_label: str = "二市場") -> list[Obs]:
    """Net purchase value (差引, 億円) per investor type for one weekly file.

    wanted maps a label keyword (e.g. '海外投資家') to a series_id.
    """
    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(xlsx)).worksheets[0]
    unit_row = next((r for r in range(1, 15)
                     if any("差引" in str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))), None)
    value_row = next((r + 1 for r in range(unit_row or 1, ws.max_row + 1)
                      if market_label in str(ws.cell(r, 2).value or "")), None)
    if unit_row is None or value_row is None:
        raise ProviderError("JPX flows: layout not recognised")
    out, seen = [], set()
    for col in range(1, ws.max_column + 1):
        if "差引" not in str(ws.cell(unit_row, col).value or ""):
            continue
        start = col - 2  # each investor block is 売 / 買 / 差引 / 合計
        labels = " > ".join(dict.fromkeys(t for r in range(2, unit_row)
                                           if (t := _cell_text(ws, r, start))))
        for keyword, series_id in wanted.items():
            if keyword in labels and series_id not in seen:
                number = _num(ws.cell(value_row, col).value)
                if number is not None:
                    seen.add(series_id)
                    out.append(Obs(series_id, week_end, round(number / THOUSAND_YEN_PER_OKU, 1)))
    return out


def jpx_weekly_links(index_html: str, limit: int) -> list[tuple[str, str]]:
    """[(absolute_url, week_end YYYY-MM-DD)] newest first."""
    found = []
    for href, _start, end in JPX_WEEKLY_FILE.findall(index_html):
        url = href if href.startswith("http") else JPX_ORIGIN + href
        item = (url, f"{end[:4]}-{end[4:6]}-{end[6:]}")
        if item not in found:
            found.append(item)
    return sorted(found, key=lambda x: x[1], reverse=True)[:limit]


# ---------------------------------------------------------------- fetchers

Getter = Callable[..., object]  # common.http_get signature


def fetch(entry: dict, start: str, get: Getter) -> list[Obs]:
    """Fetch one catalog series. start is YYYY-MM-DD (oldest date wanted)."""
    sid, provider, p = entry["id"], entry["provider"], entry.get("params", {})
    if provider == "boj":
        resp = get(BOJ_API, params={"format": "json", "lang": "jp", "db": p["db"],
                                    "code": p["code"], "startDate": start[:4] + start[5:7]})
        return parse_boj(resp.json(), sid)
    if provider == "dashboard":
        # RegionCode keeps large indicators under the API's 100,000-record limit
        resp = get(DASHBOARD_API, params={"Lang": "JP", "IndicatorCode": p["indicator"], "RegionCode": "00000"})
        obs = parse_dashboard(resp.json(), sid, p.get("cycle", "1"), p.get("seasonal"))
        return [o for o in obs if o.date >= start]
    if provider == "mof_jgb":
        obs: dict[str, Obs] = {}
        for url in MOF_JGB_URLS:  # history file + current month (the latter is fresher)
            text = get(url).content.decode("shift_jis", errors="replace")
            obs.update({o.date: o for o in parse_mof_jgb(text, sid, p["tenor"])})
        return [obs[d] for d in sorted(obs) if d >= start]
    if provider == "fred":
        key = os.environ.get("FRED_API_KEY", "")
        if not key:
            raise ProviderError("FRED_API_KEY 未設定（.env に設定すると取得できる）")
        resp = get(FRED_API, params={"series_id": p["id"], "api_key": key, "file_type": "json",
                                     "observation_start": start})
        return parse_fred(resp.json(), sid)
    if provider == "cftc":
        weeks = (dt.date.today() - dt.date.fromisoformat(start)).days // 7 + 5  # one report a week
        resp = get(CFTC_API, params={"cftc_contract_market_code": p["contract"],
                                     "$order": "report_date_as_yyyy_mm_dd DESC", "$limit": max(60, weeks),
                                     "$where": f"report_date_as_yyyy_mm_dd >= '{start}T00:00:00'"})
        return parse_cftc(resp.json(), sid)
    if provider == "yahoo":
        resp = get(YAHOO_CHART.format(symbol=p["symbol"]),
                   params={"period1": epoch(start), "period2": epoch(dt.date.today().isoformat()) + 86400,
                           "interval": "1d", "events": "history"})
        return [o for o in parse_yahoo(resp.json(), sid) if o.date >= start]
    raise ProviderError(f"unknown provider {provider}")


def fetch_jpx_flows(entries: list[dict], weeks: int, get: Getter) -> list[Obs]:
    """All jpx_flows series share the same weekly files, so they are fetched together."""
    wanted = {e["params"]["investor"]: e["id"] for e in entries}
    html = get(JPX_FLOWS_PAGE).content.decode("utf-8", errors="replace")
    out: list[Obs] = []
    for url, week_end in jpx_weekly_links(html, weeks):
        out.extend(parse_jpx_flows(get(url, referer=JPX_FLOWS_PAGE).content, week_end, wanted))
    return out


def lookback_start(years: int, today: dt.date | None = None) -> str:
    today = today or dt.date.today()
    return today.replace(year=today.year - years, month=1, day=1).isoformat()
